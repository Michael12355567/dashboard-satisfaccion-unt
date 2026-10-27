"""Motor local de series institucionales y reportes. No depende de servicios externos."""
from pathlib import Path
from io import BytesIO
import numpy as np
import pandas as pd
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, TableStyle, Spacer, Image, PageBreak, KeepTogether
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import cm
import re

BASE=Path(__file__).resolve().parent
NAVY='#092653'; GOLD='#C79600'


def leer_academico():
    df=pd.read_csv(BASE/'data'/'academico.csv',encoding='utf-8-sig')
    df['anio']=df['anio'].astype(int);df['valor']=pd.to_numeric(df['valor'])
    return df


def filtrar(df, facultad='Toda la UNT', carrera='Todas', sede='Todas', anos=None, indicadores=None):
    d=df.copy()
    if indicadores:d=d[d.indicador.isin(indicadores)]
    if anos:d=d[d.anio.isin(anos)]
    if facultad!='Toda la UNT':
        catalogo=df[(df.dimension=='Facultad') & (df.grupo==facultad)].programa.unique()
        d=d[d.programa.isin(catalogo)]
    if carrera!='Todas':d=d[d.programa==carrera]
    if sede!='Todas':
        programs=df[(df.dimension=='Sede') & (df.grupo==sede)].programa.unique()
        d=d[d.programa.isin(programs)]
        d=d[(d.dimension=='Sede') & (d.grupo==sede)]
    else:d=d[d.dimension=='Facultad']
    return d


def resumen(df):
    if df.empty:return pd.DataFrame(columns=['indicador','anio','valor'])
    return df.groupby(['indicador','anio'],as_index=False).valor.sum()


def validar(df):
    assert df[['indicador','dimension','grupo','programa','anio']].duplicated().sum()==0,'Duplicados en claves'
    assert (df.valor>=0).all(),'Valores negativos'
    return True


def tendencia(data, horizon=3):
    """Proyección lineal ilustrativa. Mín. tres observaciones; no utilizar como meta oficial."""
    data=data.dropna(subset=['anio','valor']).sort_values('anio')
    if len(data)<3:return pd.DataFrame(columns=['anio','estimacion','inferior','superior'])
    years=data.anio.to_numpy(dtype=float)
    vals=data.valor.to_numpy(dtype=float)
    co=np.polyfit(years-years.min(),vals,1)
    predicted=np.polyval(co,years-years.min())
    residual=float(np.std(vals-predicted,ddof=1)) if len(vals)>2 else 0
    fut=np.arange(int(years.max())+1,int(years.max())+horizon+1)
    yhat=np.maximum(0,np.polyval(co,fut-years.min()))
    # Margen descriptivo de incertidumbre, no intervalo de confianza formal.
    margin=1.96*residual*np.sqrt(1+np.arange(1,horizon+1)/len(vals))
    return pd.DataFrame({'anio':fut,'estimacion':np.round(yhat).astype(int),'inferior':np.maximum(0,np.round(yhat-margin)).astype(int),'superior':np.round(yhat+margin).astype(int)})


def _cover_doc(d,titulo):
    p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run('UNIVERSIDAD NACIONAL DE TRUJILLO');r.bold=True;r.font.size=Pt(15);r.font.color.rgb=RGBColor(9,38,83)
    p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run('OFICINA DE PLANEAMIENTO Y PRESUPUESTO  |  UNIDAD DE ESTADÍSTICA');r.font.size=Pt(9)
    photo=BASE/'assets'/'portada_unt.jpg'
    if photo.exists():
        p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run().add_picture(str(photo),width=Cm(14.7))
    p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run(titulo.upper());r.bold=True;r.font.size=Pt(20);r.font.color.rgb=RGBColor(9,38,83)
    p=d.add_paragraph('Serie estadística institucional · Elaboración automática con información disponible');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    d.add_page_break()


def _shade(cell,color):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd');shd.set(qn('w:fill'),color);tcPr.append(shd)


def _table_word(doc,headers,records):
    t=doc.add_table(rows=1,cols=len(headers));t.style='Table Grid';t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i];c.text=str(h);_shade(c,'092653')
        for run in c.paragraphs[0].runs:run.font.color.rgb=RGBColor(255,255,255);run.bold=True;run.font.size=Pt(8)
    for k,record in enumerate(records):
        r=t.add_row()
        for i,v in enumerate(record):
            c=r.cells[i];c.text=str(v)
            if k%2:_shade(c,'F0F4F9')
            for p in c.paragraphs:
                for run in p.runs:run.font.size=Pt(8)
    doc.add_paragraph()


def word_informe(df,titulo,scope,notes=''):
    doc=Document();section=doc.sections[0];section.top_margin=Cm(2.2);section.bottom_margin=Cm(2)
    _cover_doc(doc,titulo)
    doc.add_heading('PRESENTACIÓN',0)
    doc.add_paragraph(f'El presente informe describe los indicadores disponibles para {scope}, organizados por año e indicador. Las cifras proceden del catálogo estadístico institucional UNT 2021–2026. Se muestran únicamente las series seleccionadas y verificables dentro de la base incorporada.')
    doc.add_heading('METODOLOGÍA Y CRITERIOS ESTADÍSTICOS',1)
    doc.add_paragraph('Se presentan registros por año académico. Los datos no disponibles se representan mediante un guion (—), sin equipararlos a cero. Los totales por facultad corresponden a la suma de sus carreras. La matrícula refleja el periodo seleccionado para cada año, y los conteos de postulaciones pueden incluir más de un proceso.')
    if notes:doc.add_paragraph(notes)
    doc.add_heading('INDICADORES ESTADÍSTICOS',1)
    if df.empty:doc.add_paragraph('No hay registros para los filtros seleccionados.')
    for i,(indicador,g) in enumerate(df.groupby('indicador',sort=False),1):
        doc.add_heading(f'{i}. {indicador}',level=2)
        pv=g.groupby('anio',as_index=False).valor.sum().sort_values('anio')
        doc.add_paragraph(f'Tabla {i}',style=None).runs[0].bold=True
        p=doc.add_paragraph(f'Evolución anual de {indicador.lower()}, {int(pv.anio.min())}–{int(pv.anio.max())}');p.runs[0].italic=True
        _table_word(doc,['Año','Total'],[(int(r.anio),f'{int(r.valor):,}'.replace(',',' ')) for r in pv.itertuples()])
        doc.add_paragraph('Nota. Información estadística de la Universidad Nacional de Trujillo.').runs[0].italic=True
        if g.programa.nunique()>1:
            pv2=g.pivot_table(index='programa',columns='anio',values='valor',aggfunc='sum').sort_index()
            headers=['Carrera profesional']+[str(x) for x in pv2.columns]
            records=[]
            for career,row in pv2.iterrows():records.append([career]+['—' if pd.isna(v) else f'{int(v):,}'.replace(',',' ') for v in row.tolist()])
            doc.add_paragraph(f'Tabla {i}.1. Desagregación por carrera',style=None)
            _table_word(doc,headers,records)
    doc.add_heading('ALCANCE',1)
    doc.add_paragraph('Los resultados corresponden únicamente a los indicadores y periodos disponibles en el repositorio incorporado. Este reporte no constituye una certificación de cierre de los registros administrativos.')
    bio=BytesIO();doc.save(bio);return bio.getvalue()


def pdf_informe(df,titulo,scope):
    bio=BytesIO()
    styles=getSampleStyleSheet()
    styles.add(ParagraphStyle(name='UNTTitle',parent=styles['Title'],fontSize=17,textColor=colors.HexColor(NAVY),alignment=TA_CENTER,spaceAfter=12))
    styles.add(ParagraphStyle(name='UNTH2',parent=styles['Heading2'],textColor=colors.HexColor(NAVY),spaceBefore=12))
    story=[]
    cover=BASE/'assets'/'portada_unt.jpg'
    story.append(Paragraph('UNIVERSIDAD NACIONAL DE TRUJILLO',styles['UNTTitle']))
    if cover.exists():story.append(Image(str(cover),width=15*cm,height=7*cm,kind='proportional'))
    story.append(Spacer(1,1*cm));story.append(Paragraph(titulo,styles['UNTTitle']))
    story.append(Paragraph('Oficina de Planeamiento y Presupuesto · Unidad de Estadística',styles['Normal']))
    story.append(PageBreak())
    story.append(Paragraph('PRESENTACIÓN',styles['Heading1']))
    story.append(Paragraph(f'Informe de indicadores estadísticos para {scope}. Los resultados reflejan únicamente los registros disponibles en el catálogo institucional.',styles['Normal']))
    for i,(ind,g) in enumerate(df.groupby('indicador',sort=False),1):
        s=g.groupby('anio',as_index=False).valor.sum().sort_values('anio')
        story.append(Paragraph(f'{i}. {ind}',styles['UNTH2']))
        dat=[['Año','Total']]+[[str(int(r.anio)),f'{int(r.valor):,}'.replace(',',' ')] for r in s.itertuples()]
        t=Table(dat,colWidths=[5*cm,5*cm],repeatRows=1)
        t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor(NAVY)),('TEXTCOLOR',(0,0),(-1,0),colors.white),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#F0F4F9')]),('GRID',(0,0),(-1,-1),.35,colors.HexColor('#CED8E3')),('ALIGN',(1,1),(-1,-1),'RIGHT'),('PADDING',(0,0),(-1,-1),6)]))
        story.append(t)
        story.append(Spacer(1,.25*cm))
        story.append(Paragraph('Nota. Información estadística institucional.',styles['Italic']))
    doc=SimpleDocTemplate(bio,pagesize=A4,leftMargin=2*cm,rightMargin=2*cm,topMargin=1.8*cm,bottomMargin=1.8*cm)
    doc.build(story)
    return bio.getvalue()
