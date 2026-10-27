"""Transformaciones y Excel sin exponer datos en registros de depuración."""
from __future__ import annotations
from io import BytesIO
from datetime import datetime, timezone
import pandas as pd

NULOS_POR_ORIGEN = {"discapacitado", "tipo_discapacidad", "escala_pago", "movilidad", "identidad_etnica"}
MAX_EXCEL_FILAS = 150_000


def tabla(filas: list[dict]) -> pd.DataFrame:
    if not filas:
        return pd.DataFrame()
    return pd.json_normalize(filas, sep="_")


def calidad(df: pd.DataFrame, advertencias=None) -> dict:
    filas = len(df)
    campos = len(df.columns)
    afectados = []
    for campo in df.columns:
        if campo not in NULOS_POR_ORIGEN:
            cuenta = int(df[campo].isna().sum())
            if df[campo].dtype == "object":
                cuenta += int((df[campo].fillna("").astype(str).str.strip().eq("")).sum()) - int(df[campo].isna().sum())
            if cuenta:
                afectados.append({"campo": campo, "vacios": cuenta, "porcentaje": round(100 * cuenta / max(filas, 1), 2)})
    dni_invalidos = 0
    if {"tipo_documento", "nro_documento"}.issubset(df.columns):
        es_dni = df["tipo_documento"].fillna("").astype(str).str.upper().eq("DNI")
        valido = df["nro_documento"].fillna("").astype(str).str.fullmatch(r"\d{8}")
        dni_invalidos = int((es_dni & ~valido).sum())
    avisos = advertencias or []
    avisos_resumen = {}
    for a in avisos:
        if isinstance(a, dict):
            k = (str(a.get("campo", "sin_campo")), str(a.get("codigo", "sin_codigo")))
            avisos_resumen[k] = avisos_resumen.get(k, 0) + 1
    return {"filas": filas, "columnas": campos, "dni_invalidos": dni_invalidos,
            "nulos": pd.DataFrame(afectados),
            "advertencias": pd.DataFrame([{"campo": k[0], "codigo": k[1], "cantidad": n}
                                         for k, n in avisos_resumen.items()])}


def resumen(df: pd.DataFrame) -> pd.DataFrame:
    # Agrupaciones descriptivas de filas, NO personas únicas o cifras oficiales.
    if df.empty:
        return pd.DataFrame(columns=["dimension", "categoria", "filas"])
    agrupaciones = []
    for campo in ("periodo_matricula", "periodo_egreso", "carrera", "local", "genero", "modalidad_estudios", "Des_sede", "Nom_escuela"):
        if campo in df.columns:
            conteos = df[campo].fillna("Sin dato").astype(str).value_counts(dropna=False)
            agrupaciones.extend({"dimension": campo, "categoria": nombre, "filas": int(n)}
                                for nombre, n in conteos.items())
    return pd.DataFrame(agrupaciones, columns=["dimension", "categoria", "filas"])


def _sin_formula(x):
    # Bloquea fórmulas que pudieran ejecutarse al abrir Excel. Solo Excel; no altera el JSON fuente.
    if isinstance(x, str) and x.lstrip().startswith(("=", "+", "-", "@", "\t", "\r")):
        return "'" + x
    return x


def excel(filas: list[dict], advertencias: list[dict], metadatos: dict) -> bytes:
    df = tabla(filas)
    if len(df) > MAX_EXCEL_FILAS:
        raise ValueError("El Excel supera 150 000 registros. Solicita menos períodos para evitar agotamiento de memoria.")
    qc = calidad(df, advertencias)
    data = BytesIO()
    hojas = {
        "Datos": df,
        "Resumen_filas": resumen(df),
        "Calidad_nulos": qc["nulos"],
        "Advertencias_resumen": qc["advertencias"],
        "Advertencias_originales": tabla(advertencias),
        "Metadatos": pd.DataFrame([{"campo": str(k), "valor": str(v)} for k, v in metadatos.items()]),
    }
    with pd.ExcelWriter(data, engine="openpyxl") as writer:
        for nombre, marco in hojas.items():
            limpio = marco.map(_sin_formula) if not marco.empty else marco
            limpio.to_excel(writer, sheet_name=nombre, index=False)
            ws = writer.sheets[nombre]
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for celda in ws[1]:
                from openpyxl.styles import Font, PatternFill
                celda.fill = PatternFill("solid", fgColor="163A63")
                celda.font = Font(color="FFFFFF", bold=True)
            for col in ws.columns:
                letra = col[0].column_letter
                ancho = max((len(str(c.value)) for c in list(col)[:150] if c.value is not None), default=10)
                ws.column_dimensions[letra].width = min(max(ancho + 2, 12), 42)
            # Los identificadores no deben perder ceros iniciales.
            for cab in ws[1]:
                if str(cab.value) in ("nro_documento", "Nro_documento", "Nom_documento", "Cod_alumno"):
                    for row in ws.iter_rows(min_row=2, min_col=cab.column, max_col=cab.column):
                        row[0].number_format = "@"
    return data.getvalue()
